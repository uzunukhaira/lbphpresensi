from io import StringIO
import csv
from flask import Blueprint, Response, request, jsonify, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from database import get_connection, release_connection
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from datetime import datetime

admin_bp = Blueprint('admin_bp', __name__)

@admin_bp.route('/admin/presensi', methods=['GET'])
def get_admin_presensi():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        
        # PERBAIKAN: Melakukan JOIN lengkap ke tabel kelas, jadwal, dan matakuliah
        cursor.execute("""
            SELECT p.id_presensi, p.nim, m.nama_lengkap, k.nama_kelas, mk.nama_mk, p.status_absen, p.tanggal, p.waktu_absen 
            FROM presensi p
            JOIN mahasiswa m ON p.nim = m.nim
            LEFT JOIN kelas k ON m.id_kelas = k.id_kelas
            LEFT JOIN jadwal_kuliah jk ON p.id_jadwal = jk.id_jadwal
            LEFT JOIN matakuliah mk ON jk.id_mk = mk.id_mk
            ORDER BY p.tanggal DESC, p.waktu_absen DESC LIMIT 50
        """)
        
        rows = cursor.fetchall()
        rekap_list = []
        for r in rows:
            # Menggabungkan tanggal dan waktu dengan aman
            tgl_str = str(r[6]) if r[6] else ""
            waktu_str = str(r[7]) if r[7] else ""
            waktu_gabung = f"{tgl_str} {waktu_str}".strip()

            rekap_list.append({
                "id": r[0], 
                "nim": r[1], 
                "nama": r[2], 
                "nama_kelas": r[3] if r[3] else "-", # Nama Kelas
                "mata_kuliah": r[4] if r[4] else "Tidak Diketahui", # Mata Kuliah
                "status": r[5],
                "waktu": waktu_gabung if waktu_gabung else "-"
            })
            
        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "data": rekap_list}), 200
        
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@admin_bp.route('/admin/mahasiswa', methods=['GET'])
def admin_get_mahasiswa():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        
        # PERBAIKAN: JOIN ke tabel kelas agar nama_kelas dan id_kelas bisa tampil di web
        cursor.execute("""
            SELECT m.nim, m.nama_lengkap, m.wajah_terdaftar, m.id_kelas, k.nama_kelas 
            FROM mahasiswa m 
            LEFT JOIN kelas k ON m.id_kelas = k.id_kelas 
            ORDER BY m.nim ASC
        """)
        
        rows = cursor.fetchall()
        cursor.close()
        release_connection(conn)
        
        data = [{
            "nim": r[0], 
            "nama": r[1], 
            "wajah_terdaftar": r[2],
            "id_kelas": r[3],
            "nama_kelas": r[4]
        } for r in rows]
        
        return jsonify({"status": "success", "data": data}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/mahasiswa/<nim>', methods=['PUT'])
def admin_edit_mahasiswa(nim):
    nama = request.form.get('nama')
    id_kelas = request.form.get('id_kelas')
    password = request.form.get('password')

    if not nama or not id_kelas:
        return jsonify({"status": "error", "message": "Nama dan Kelas wajib diisi!"}), 400

    try:
        conn = get_connection()
        cursor = conn.cursor()

        # Logika: Jika admin mengetik password baru, update passwordnya
        if password and password.strip() != "":
            hashed_pw = generate_password_hash(password)
            cursor.execute("""
                UPDATE mahasiswa 
                SET nama_lengkap = %s, id_kelas = %s, password = %s 
                WHERE nim = %s
            """, (nama, id_kelas, hashed_pw, nim))
        # Jika password di form kosong, jangan ubah password di database
        else:
            cursor.execute("""
                UPDATE mahasiswa 
                SET nama_lengkap = %s, id_kelas = %s 
                WHERE nim = %s
            """, (nama, id_kelas, nim))

        conn.commit()
        cursor.close()
        release_connection(conn)

        return jsonify({"status": "success", "message": f"Data Mahasiswa {nim} berhasil diperbarui!"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/mahasiswa/<nim>', methods=['DELETE'])
def admin_delete_mahasiswa(nim):
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM mahasiswa WHERE nim = %s", (nim,))
        conn.commit()
        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "message": f"Mahasiswa {nim} berhasil dihapus!"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/master-data', methods=['GET'])
def get_master_data():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id_kelas, nama_kelas FROM kelas")
        kelas = [{"id": r[0], "nama": r[1]} for r in cursor.fetchall()]
        cursor.execute("SELECT id_dosen, nama_dosen FROM dosen")
        dosen = [{"id": r[0], "nama": r[1]} for r in cursor.fetchall()]
        cursor.execute("SELECT id_mk, nama_mk FROM matakuliah")
        matkul = [{"id": r[0], "nama": r[1]} for r in cursor.fetchall()]
        cursor.execute("SELECT id_ruangan, nama_ruangan FROM ruangan")
        ruangan = [{"id": r[0], "nama": r[1]} for r in cursor.fetchall()]
        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "data": {"kelas": kelas, "dosen": dosen, "matkul": matkul, "ruangan": ruangan}}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@admin_bp.route('/admin/kelas', methods=['GET'])
def admin_get_kelas():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id_kelas, nama_kelas FROM kelas ORDER BY nama_kelas ASC")
        rows = cursor.fetchall()
        cursor.close()
        release_connection(conn)
        
        data = [{"id_kelas": r[0], "nama_kelas": r[1]} for r in rows]
        return jsonify({"status": "success", "data": data}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/jadwal', methods=['GET', 'POST'])
def kelola_jadwal():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        if request.method == 'GET':
            query = """
                SELECT j.id_jadwal, m.nama_mk, d.nama_dosen, k.nama_kelas, r.nama_ruangan, 
                       j.hari, j.jam_mulai, j.jam_selesai, j.tahun_akademik, j.semester
                FROM jadwal_kuliah j
                JOIN matakuliah m ON j.id_mk = m.id_mk
                JOIN dosen d ON j.id_dosen = d.id_dosen
                JOIN kelas k ON j.id_kelas = k.id_kelas
                LEFT JOIN ruangan r ON j.id_ruangan = r.id_ruangan
                ORDER BY j.hari, j.jam_mulai
            """
            cursor.execute(query)
            rows = cursor.fetchall()
            hasil = [{
                "id_jadwal": r[0], "nama_mk": r[1], "nama_dosen": r[2], "nama_kelas": r[3],
                "nama_ruangan": r[4], "hari": r[5], "jam_mulai": str(r[6]), "jam_selesai": str(r[7]),
                "tahun_akademik": r[8], "semester": r[9]
            } for r in rows]
            cursor.close()
            release_connection(conn)
            return jsonify({"status": "success", "data": hasil}), 200

        elif request.method == 'POST':
            data = request.json
            id_mk = data.get('id_mk')
            id_dosen = data.get('id_dosen')
            id_kelas = data.get('id_kelas')
            id_ruangan = data.get('id_ruangan')
            hari = data.get('hari')
            jam_mulai = data.get('jam_mulai')
            jam_selesai = data.get('jam_selesai')
            tahun_akademik = data.get('tahun_akademik', '2023/2024')
            semester = data.get('semester', 'Genap')

            if not all([id_mk, id_dosen, id_kelas, hari, jam_mulai, jam_selesai]):
                return jsonify({"status": "error", "message": "Semua kolom wajib diisi!"}), 400

            cursor.execute("""
                INSERT INTO jadwal_kuliah (id_mk, id_dosen, id_kelas, id_ruangan, hari, jam_mulai, jam_selesai, tahun_akademik, semester)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (int(id_mk), int(id_dosen), int(id_kelas), int(id_ruangan) if id_ruangan else None, hari, jam_mulai, jam_selesai, tahun_akademik, semester))
            
            conn.commit()
            cursor.close()
            release_connection(conn)
            return jsonify({"status": "success", "message": "Jadwal berhasil ditambahkan"}), 200

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/jadwal/<int:id_jadwal>', methods=['DELETE'])
def admin_delete_jadwal(id_jadwal):
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM jadwal_kuliah WHERE id_jadwal = %s", (id_jadwal,))
        conn.commit()
        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "message": "Jadwal berhasil dihapus!"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/login', methods=['POST'])
def admin_login():
    data = request.json or {}
    username, password = data.get('username'), data.get('password')
    if not username or not password:
        return jsonify({"status": "error", "message": "Username dan Password wajib diisi!"}), 400

    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id_admin, username, nama_admin, password FROM admin WHERE username = %s", (username,))
        row = cursor.fetchone()
        cursor.close()
        release_connection(conn)

        if not row or not check_password_hash(row[3], password):
            return jsonify({"status": "error", "message": "Username atau Password Admin salah!"}), 401

        return jsonify({
            "status": "success", "message": f"Welcome, {row[2]}!",
            "token": f"admin-session-{row[0]}",
            "data": {"id": row[0], "username": row[1], "nama": row[2]}
        }), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@admin_bp.route('/admin/register', methods=['POST'])
def admin_register():
    data = request.json or {}
    username, nama, password = data.get('username'), data.get('nama'), data.get('password')
    if not username or not nama or not password:
        return jsonify({"status": "error", "message": "Semua data wajib diisi!"}), 400

    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO admin (username, nama_admin, password) VALUES (%s, %s, %s) ON CONFLICT (username) DO UPDATE SET password = EXCLUDED.password, nama_admin = EXCLUDED.nama_admin",
            (username, nama, generate_password_hash(password))
        )
        conn.commit()
        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "message": f"Admin '{username}' berhasil disimpan!"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@admin_bp.route('/admin/export/rekap', methods=['GET'])
def export_rekap_kelas():
    id_kelas = request.args.get('id_kelas')
    id_jadwal = request.args.get('id_jadwal')

    if not id_kelas or not id_jadwal:
        return jsonify({"status": "error", "message": "Parameter id_kelas dan id_jadwal wajib diisi"}), 400

    try:
        conn = get_connection()
        cursor = conn.cursor()

        # 1. Ambil info Jadwal, Matkul, Kelas, dan Dosen
        cursor.execute("""
            SELECT mk.nama_mk, k.nama_kelas, d.nama_dosen 
            FROM jadwal_kuliah jk
            JOIN matakuliah mk ON jk.id_mk = mk.id_mk
            JOIN kelas k ON k.id_kelas = %s
            JOIN dosen d ON jk.id_dosen = d.id_dosen
            WHERE jk.id_jadwal = %s
        """, (id_kelas, id_jadwal))
        info = cursor.fetchone()
        
        if not info:
            return jsonify({"status": "error", "message": "Jadwal atau Kelas tidak ditemukan"}), 404
            
        nama_mk, nama_kelas, nama_dosen = info

        # 2. Ambil data kehadiran mahasiswa
        query = """
            SELECT m.nim, m.nama_lengkap,
                   COUNT(CASE WHEN p.status_absen IN ('Hadir', 'Terlambat') THEN 1 END) as hadir,
                   COUNT(CASE WHEN p.status_absen IN ('Izin', 'Sakit') THEN 1 END) as izin,
                   COUNT(CASE WHEN p.status_absen = 'Alpha' THEN 1 END) as alpha
            FROM mahasiswa m
            JOIN jadwal_kuliah jk ON jk.id_jadwal = %s
            LEFT JOIN presensi p ON m.nim = p.nim AND p.id_jadwal = jk.id_jadwal
            WHERE m.id_kelas = %s
            GROUP BY m.nim, m.nama_lengkap
            ORDER BY m.nim ASC
        """
        cursor.execute(query, (id_jadwal, id_kelas))
        rows = cursor.fetchall()

        cursor.close()
        release_connection(conn)

        # 3. Buat Workbook Excel menggunakan openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rekap Presensi"

        # Tampilkan garis grid agar rapi
        ws.views.sheetView[0].showGridLines = True

        # Style & Font
        font_title = Font(name="Calibri", size=12, bold=True)
        font_subtitle = Font(name="Calibri", size=11, bold=True)
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_sub_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_body = Font(name="Calibri", size=10)
        font_footer = Font(name="Calibri", size=10)

        fill_header = PatternFill(start_color="366092", end_color="366092", fill_type="solid") # Biru Profesional
        fill_sub = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )

        # Bagian Kop / Judul
        ws['C1'] = "DAFTAR REKAPITULASI PRESENSI MAHASISWA"
        ws['C1'].font = font_title
        ws['C2'] = "PROGRAM STUDI TEKNOLOGI REKAYASA PERANGKAT LUNAK"
        ws['C2'].font = font_subtitle

        # Identitas Kelas & Matkul
        ws['A4'] = "KELAS"
        ws['B4'] = ":"
        ws['C4'] = nama_kelas

        ws['A5'] = "MATA KULIAH"
        ws['B5'] = ":"
        ws['C5'] = nama_mk

        ws['A6'] = "DOSEN"
        ws['B6'] = ":"
        ws['C6'] = nama_dosen

        for r in range(4, 7):
            ws[f'A{r}'].font = font_subtitle
            ws[f'C{r}'].font = font_body

        # Header Tabel Bertingkat (Mirip Gambar Template)
        ws.merge_cells('A8:A9')
        ws['A8'] = "NO"
        ws.merge_cells('B8:B9')
        ws['B8'] = "NO. BP"
        ws.merge_cells('C8:C9')
        ws['C8'] = "NAMA"

        ws.merge_cells('D8:E8')
        ws['D8'] = "JUMLAH"
        ws['D9'] = "IZIN"
        ws['E9'] = "ALFA"

        ws.merge_cells('F8:F9')
        ws['F8'] = "HADIR"
        ws.merge_cells('G8:G9')
        ws['G8'] = "PERSENTASE"

        # Styling Header Tabel
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            cell1 = ws[f'{col}8']
            cell1.font = font_header
            cell1.fill = fill_header
            cell1.alignment = Alignment(horizontal='center', vertical='center')
            cell1.border = thin_border
            
            cell2 = ws[f'{col}9']
            if cell2.value is not None:
                cell2.font = font_sub_header
                cell2.fill = fill_sub
                cell2.alignment = Alignment(horizontal='center', vertical='center')
                cell2.border = thin_border

        ws['D8'].alignment = Alignment(horizontal='center', vertical='center')
        ws['E8'].alignment = Alignment(horizontal='center', vertical='center')
        ws['D8'].fill = fill_header
        ws['E8'].fill = fill_header
        ws['D8'].font = font_header
        ws['E8'].font = font_header
        ws['D8'].border = thin_border
        ws['E8'].border = thin_border

        # Masukkan Data Mahasiswa
        row_idx = 10
        no = 1
        for r in rows:
            nim, nama, hadir, izin, alpha = r
            total_berjalan = hadir + izin + alpha
            persentase = round((hadir + izin) / total_berjalan * 100, 2) if total_berjalan > 0 else 0

            ws.cell(row=row_idx, column=1, value=no).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=2, value=str(nim)).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=3, value=nama).alignment = Alignment(horizontal='left')
            ws.cell(row=row_idx, column=4, value=izin).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=5, value=alpha).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=6, value=hadir).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=7, value=f"{persentase}%").alignment = Alignment(horizontal='center')

            for c in range(1, 8):
                cell = ws.cell(row=row_idx, column=c)
                cell.font = font_body
                cell.border = thin_border
            
            row_idx += 1
            no += 1

        # Bagian Tanda Tangan di Bawah
        row_idx += 2
        months = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
        now = datetime.now()
        tgl_indo = f"{now.day:02d} {months[now.month]} {now.year}"

        ws.cell(row=row_idx, column=6, value=f"Padang, {tgl_indo}").font = font_footer
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="Mengetahui,").font = font_footer
        ws.cell(row=row_idx, column=6, value="Dosen Pengampu").font = font_footer
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="Ketua Jurusan Teknologi Informasi").font = font_footer
        row_idx += 4
        ws.cell(row=row_idx, column=2, value="( Dr. Ir. Yuhefizar, S.Kom., M.Kom. )").font = font_footer
        ws.cell(row=row_idx, column=6, value=nama_dosen).font = font_footer
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="NIP. 19760113 200604 1 002").font = font_footer

        # Otomatis sesuaikan lebar kolom
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Simpan ke memori BytesIO agar bisa langsung diunduh browser
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        nama_file = f"Rekap_{nama_kelas}_{nama_mk.replace(' ', '_')}.xlsx"

        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=nama_file
        )

    except Exception as e:
        return jsonify({"status": "error", "message": f"Gagal Export Excel: {str(e)}"}), 500

@admin_bp.route('/admin/export/rekap-umum', methods=['GET'])
def export_rekap_umum():
    id_kelas = request.args.get('id_kelas')

    if not id_kelas:
        return jsonify({"status": "error", "message": "Parameter id_kelas wajib diisi"}), 400

    try:
        conn = get_connection()
        cursor = conn.cursor()

        # Ambil nama kelas
        cursor.execute("SELECT nama_kelas FROM kelas WHERE id_kelas = %s", (id_kelas,))
        kelas_row = cursor.fetchone()
        
        if not kelas_row:
            return jsonify({"status": "error", "message": "Kelas tidak ditemukan"}), 404
            
        nama_kelas = kelas_row[0]

        # Ambil data rekap umum (akumulasi semua matakuliah)
        query = """
            SELECT m.nim, m.nama_lengkap,
                   COUNT(CASE WHEN p.status_absen = 'Hadir' THEN 1 END) as hadir,
                   COUNT(CASE WHEN p.status_absen = 'Izin' THEN 1 END) as izin,
                   COUNT(CASE WHEN p.status_absen = 'Sakit' THEN 1 END) as sakit,
                   COUNT(CASE WHEN p.status_absen = 'Alpha' THEN 1 END) as alpha
            FROM mahasiswa m
            LEFT JOIN presensi p ON m.nim = p.nim
            WHERE m.id_kelas = %s
            GROUP BY m.nim, m.nama_lengkap
            ORDER BY m.nim ASC
        """
        cursor.execute(query, (id_kelas,))
        rows = cursor.fetchall()

        cursor.close()
        release_connection(conn)

        # --- MULAI PROSES BUAT EXCEL ---
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Rekap Umum {nama_kelas}"

        # Tampilkan garis grid agar rapi
        ws.views.sheetView[0].showGridLines = True

        # Style & Font
        font_title = Font(name="Calibri", size=12, bold=True)
        font_subtitle = Font(name="Calibri", size=11, bold=True)
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_body = Font(name="Calibri", size=10)
        font_footer = Font(name="Calibri", size=10)

        fill_header = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )

        # Bagian Kop / Judul
        ws['C1'] = "DAFTAR REKAPITULASI PRESENSI UMUM KELAS"
        ws['C1'].font = font_title
        ws['C2'] = "PROGRAM STUDI TEKNOLOGI REKAYASA PERANGKAT LUNAK"
        ws['C2'].font = font_subtitle

        # Identitas Kelas
        ws['A4'] = "KELAS"
        ws['B4'] = ":"
        ws['C4'] = nama_kelas

        for r in range(4, 5):
            ws[f'A{r}'].font = font_subtitle
            ws[f'C{r}'].font = font_body

        # Header Tabel
        headers = ["NO", "NO. BP", "NAMA MAHASISWA", "HADIR", "IZIN", "SAKIT", "ALPHA"]
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col_num, value=header_title)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        # Masukkan Data Mahasiswa
        row_idx = 7
        no = 1
        for r in rows:
            nim, nama, hadir, izin, sakit, alpha = r

            ws.cell(row=row_idx, column=1, value=no).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=2, value=str(nim)).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=3, value=nama).alignment = Alignment(horizontal='left')
            ws.cell(row=row_idx, column=4, value=hadir or 0).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=5, value=izin or 0).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=6, value=sakit or 0).alignment = Alignment(horizontal='center')
            ws.cell(row=row_idx, column=7, value=alpha or 0).alignment = Alignment(horizontal='center')

            # Berikan garis batas pada setiap sel data
            for c in range(1, 8):
                ws.cell(row=row_idx, column=c).font = font_body
                ws.cell(row=row_idx, column=c).border = thin_border
            
            row_idx += 1
            no += 1

        # Bagian Tanda Tangan di Bawah
        row_idx += 2
        months = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
        now = datetime.now()
        tgl_indo = f"{now.day:02d} {months[now.month]} {now.year}"

        ws.cell(row=row_idx, column=6, value=f"Padang, {tgl_indo}").font = font_footer
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="Mengetahui,").font = font_footer
        ws.cell(row=row_idx, column=6, value="Admin Akademik").font = font_footer
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="Ketua Jurusan Teknologi Informasi").font = font_footer
        row_idx += 4
        ws.cell(row=row_idx, column=2, value="( Dr. Ir. Yuhefizar, S.Kom., M.Kom. )").font = font_footer
        ws.cell(row=row_idx, column=6, value="(.......................................)").font = font_footer
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="NIP. 19760113 200604 1 002").font = font_footer

        # Otomatis sesuaikan lebar kolom biar nggak sempit
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Simpan ke memori untuk di-download
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        nama_file = f"Rekap_Umum_Kelas_{nama_kelas}.xlsx"

        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=nama_file
        )

    except Exception as e:
        return jsonify({"status": "error", "message": f"Gagal Export Excel: {str(e)}"}), 500

@admin_bp.route('/admin/rekap-tabel-umum', methods=['GET'])
def get_rekap_tabel_umum_json():
    id_kelas = request.args.get('id_kelas')
    if not id_kelas:
        return jsonify({"status": "error", "message": "Parameter id_kelas wajib diisi"}), 400

    try:
        conn = get_connection()
        cursor = conn.cursor()

        # Query rekap total (akumulasi semua matkul)
        query = """
            SELECT m.nim, m.nama_lengkap,
                   COUNT(CASE WHEN p.status_absen = 'Hadir' THEN 1 END) as hadir,
                   COUNT(CASE WHEN p.status_absen = 'Izin' THEN 1 END) as izin,
                   COUNT(CASE WHEN p.status_absen = 'Sakit' THEN 1 END) as sakit,
                   COUNT(CASE WHEN p.status_absen = 'Alpha' THEN 1 END) as alpha
            FROM mahasiswa m
            LEFT JOIN presensi p ON m.nim = p.nim
            WHERE m.id_kelas = %s
            GROUP BY m.nim, m.nama_lengkap
            ORDER BY m.nim ASC
        """
        cursor.execute(query, (id_kelas,))
        rows = cursor.fetchall()

        hasil = []
        for r in rows:
            nim, nama, hadir, izin, sakit, alpha = r
            hasil.append({
                "nim": nim,
                "nama": nama,
                "hadir": hadir or 0,
                "izin": izin or 0,
                "sakit": sakit or 0,
                "alpha": alpha or 0
            })

        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "data": hasil}), 200

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@admin_bp.route('/admin/izin', methods=['GET'])
def get_admin_izin():
    try:
        conn = get_connection()
        cursor = conn.cursor()
        
        # Ambil data presensi yang statusnya Izin atau Sakit
        cursor.execute("""
            SELECT p.id_presensi, p.nim, m.nama_lengkap, mk.nama_mk, p.status_absen, p.bukti_izin, p.waktu_absen 
            FROM presensi p
            JOIN mahasiswa m ON p.nim = m.nim
            LEFT JOIN jadwal_kuliah jk ON p.id_jadwal = jk.id_jadwal
            LEFT JOIN matakuliah mk ON jk.id_mk = mk.id_mk
            WHERE p.status_absen IN ('Izin', 'Sakit')
            ORDER BY p.waktu_absen DESC
        """)
        rows = cursor.fetchall()
        
        izin_list = []
        for r in rows:
            izin_list.append({
                "id": r[0],
                "nim": r[1],
                "nama": r[2],
                "mata_kuliah": r[3] if r[3] else "-",
                "status": r[4],
                "bukti_url": r[5], # Path file surat di server
                "waktu": r[6].strftime("%d-%m-%Y %H:%M:%S") if r[6] else "-"
            })
            
        cursor.close()
        release_connection(conn)
        return jsonify({"status": "success", "data": izin_list}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500